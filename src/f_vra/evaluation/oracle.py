"""Import manual reachability labels as candidates with explicit review gates."""

from __future__ import annotations

import hashlib
import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

from packaging.utils import canonicalize_name

from f_vra.evaluation.contracts import Decision, OracleLabel
from f_vra.evaluation.io import read_parquet, write_parquet


def _read_manual_annotation(source: Path) -> tuple[list[Any], list[tuple[Any, ...] | list[str]]]:
    """Read the manual annotation table from an Excel workbook or CSV export."""
    if source.suffix.lower() == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            table = list(csv.reader(handle))
        if not table:
            raise ValueError("manual annotation CSV is empty")
        return list(table[0]), table[1:]

    from openpyxl import load_workbook

    book = load_workbook(source, read_only=True, data_only=True)
    try:
        if "manual_annotation" not in book.sheetnames:
            raise ValueError("workbook does not contain a manual_annotation sheet")
        rows = book["manual_annotation"].values
        return list(next(rows)), list(rows)
    finally:
        book.close()


def import_oracle(workbook: Path, dataset_dir: Path) -> dict[str, Any]:

    candidates = read_parquet(dataset_dir / "query_candidates.parquet")
    if not candidates:
        raise ValueError("query_candidates.parquet is empty; freeze the dataset first")
    queries_path = dataset_dir / "queries.parquet"
    queries = read_parquet(queries_path) if queries_path.exists() else []
    by_key = {(str(row["project_id"]), str(row["cve_id"])): row for row in candidates}
    query_by_key = {(str(row["project_id"]), str(row["cve_id"])): row for row in queries if row.get("project_id")}
    headings, rows = _read_manual_annotation(workbook)
    columns = {str(value): index for index, value in enumerate(headings)}
    required = {"project", "cve", "package", "version", "manual_label", "manual_reason", "Path(nếu reach)", "annotator"}
    if not required.issubset(columns):
        raise ValueError(f"missing oracle columns: {sorted(required - columns.keys())}")
    source_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    labels: list[OracleLabel] = []
    audit: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for row_number, row in enumerate(rows, 2):
        def cell(name: str) -> str:
            index = columns[name]
            return str(row[index]).strip() if index < len(row) and row[index] is not None else ""

        # manual_label is the only verdict source. auto_reachability is intentionally ignored.
        raw_label = cell("manual_label").lower()
        if raw_label not in {"reachable", "unreachable", "unknown"}:
            continue
        key = (cell("project"), cell("cve"))
        if key in seen:
            raise ValueError(f"duplicate manual label: {key}")
        seen.add(key)
        candidate = by_key.get(key)
        if candidate is None:
            audit.append({"row": row_number, "project_id": key[0], "cve_id": key[1], "status": "NO_CANDIDATE"})
            continue
        names = {canonicalize_name(part.strip()) for part in cell("package").split(";") if part.strip()}
        expected = canonicalize_name(str(candidate["affected_distribution"]))
        if expected not in names:
            audit.append({
                "row": row_number, "project_id": key[0], "cve_id": key[1],
                "status": "PACKAGE_MISMATCH", "expected": expected,
            })
            continue
        versions = {part.strip() for part in cell("version").replace(",", ";").split(";") if part.strip()}
        expected_versions = {str(value) for value in candidate.get("affected_versions_claimed") or []}
        if not versions.intersection(expected_versions):
            audit.append({
                "row": row_number, "project_id": key[0], "cve_id": key[1],
                "status": "VERSION_MISMATCH", "expected": sorted(expected_versions),
            })
            continue
        reason, witness = cell("manual_reason"), cell("Path(nếu reach)")
        decision = Decision(raw_label.upper())
        review_reasons = []
        if not reason:
            review_reasons.append("MISSING_REASON")
        if decision == Decision.REACHABLE and not witness:
            review_reasons.append("MISSING_WITNESS")
        if not cell("annotator"):
            review_reasons.append("MISSING_ANNOTATOR")
        # A manual label never proves source completeness or witness validity by itself.
        review_reasons.append("SNAPSHOT_EVIDENCE_NOT_ADJUDICATED")
        labels.append(OracleLabel(
            candidate_id=str(candidate["candidate_id"]),
            query_id=str(query_by_key.get(key, {}).get("query_id") or ""),
            decision=decision, reason=reason,
            witness=witness, annotator=cell("annotator"), source_workbook_hash=source_hash,
            review_reason=";".join(review_reasons),
        ))
        audit.append({
            "row": row_number, "project_id": key[0], "cve_id": key[1], "status": "MATCHED",
            "candidate_id": candidate["candidate_id"],
            "query_id": query_by_key.get(key, {}).get("query_id") or "",
            "review_reason": ";".join(review_reasons),
        })
    write_parquet(dataset_dir / "oracle_labels.parquet", [item.model_dump(mode="json") for item in labels])
    write_parquet(dataset_dir / "oracle_import_audit.parquet", audit)
    eligibility_path = dataset_dir / "query_eligibility.parquet"
    eligibility = {
        str(row["candidate_id"]): row
        for row in (read_parquet(eligibility_path) if eligibility_path.exists() else [])
    }
    candidates_by_id = {str(row["candidate_id"]): row for row in candidates}
    review_queue = []
    for label in labels:
        row = label.model_dump(mode="json")
        candidate = candidates_by_id[label.candidate_id]
        eligible = eligibility.get(label.candidate_id, {})
        reasons = set(label.review_reason.split(";"))
        if "MISSING_REASON" in reasons or "MISSING_ANNOTATOR" in reasons:
            priority = 1
        elif "MISSING_WITNESS" in reasons:
            priority = 2
        elif eligible.get("query_eligible") and label.decision == Decision.REACHABLE:
            priority = 3
        elif eligible.get("query_eligible"):
            priority = 4
        else:
            priority = 5
        review_queue.append({
            **row,
            "project_id": candidate["project_id"], "cve_id": candidate["cve_id"],
            "affected_distribution": candidate["affected_distribution"],
            "affected_versions_claimed": candidate.get("affected_versions_claimed") or [],
            "query_eligible": bool(eligible.get("query_eligible")),
            "eligibility_status": str(eligible.get("status") or "ENVIRONMENT_NOT_AUDITED"),
            "review_priority": priority,
        })
    review_queue.sort(key=lambda item: (item["review_priority"], item["project_id"], item["cve_id"]))
    write_parquet(dataset_dir / "oracle_review_queue.parquet", review_queue)
    csv_path = dataset_dir / "oracle_review_queue.csv"
    csv_fields = [
        "review_priority", "candidate_id", "query_id", "project_id", "cve_id",
        "affected_distribution", "decision", "reason", "witness", "annotator",
        "query_eligible", "eligibility_status", "review_reason",
    ]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=csv_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(review_queue)
    with (dataset_dir / "oracle_review_queue_eligible.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=csv_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(row for row in review_queue if row["query_eligible"])
    template_path = dataset_dir / "oracle_review_template.jsonl"
    template_path.write_text("".join(
        json.dumps({
            "candidate_id": row["candidate_id"], "review_status": "pending_review",
            "evidence": "", "reviewer": "", "reviewed_at": "",
            "witness_validated": False, "frontier_complete": False,
        }, ensure_ascii=False) + "\n"
        for row in review_queue if row["query_eligible"]
    ), encoding="utf-8")
    counts = Counter(item["status"] for item in audit)
    label_counts = Counter(item.decision.value for item in labels)
    return {
        "matched": len(labels), "status_counts": dict(counts),
        "source_sha256": source_hash, "source_format": workbook.suffix.lower().lstrip("."),
        "label_column": "manual_label",
        "ignored_verdict_columns": ["auto_reachability"] if "auto_reachability" in columns else [],
        "manual_label_counts": dict(label_counts),
        "eligible_for_review": sum(bool(row["query_eligible"]) for row in review_queue),
        "pending_review": len(review_queue),
    }


def adjudicate_oracle(dataset_dir: Path, decisions: Path) -> dict[str, Any]:
    """Apply reviewer decisions without ever inferring acceptance from a workbook label."""
    import json

    rows = read_parquet(dataset_dir / "oracle_labels.parquet")
    labels = {str(row.get("candidate_id") or row.get("query_id")): row for row in rows}
    query_to_candidate = {str(row["query_id"]): key for key, row in labels.items() if row.get("query_id")}
    eligibility = {
        str(row["candidate_id"]): row
        for row in read_parquet(dataset_dir / "query_eligibility.parquet")
    }
    updates = [json.loads(line) for line in decisions.read_text(encoding="utf-8").splitlines() if line.strip()]
    accepted = rejected = 0
    for update in updates:
        label_id = str(update.get("candidate_id") or update.get("query_id") or "")
        key = label_id if label_id in labels else query_to_candidate.get(label_id, "")
        if key not in labels:
            raise ValueError(f"unknown oracle candidate/query id {label_id}")
        status = str(update["review_status"])
        if status not in {"accepted", "rejected", "pending_review"}:
            raise ValueError(f"invalid review_status {status}")
        row = labels[key]
        if status == "accepted":
            missing = [
                field for field in ("evidence", "reviewer", "reviewed_at")
                if not str(update.get(field) or "").strip()
            ]
            if missing:
                raise ValueError(f"accepted oracle {label_id} requires {', '.join(missing)}")
            if not eligibility.get(str(row["candidate_id"]), {}).get("query_eligible"):
                raise ValueError(f"accepted oracle {label_id} is not query eligible")
            if row["decision"] == Decision.REACHABLE.value:
                if not str(row.get("witness") or "").strip() or not bool(update.get("witness_validated")):
                    raise ValueError(f"accepted reachable oracle {label_id} requires a validated pinned-source witness")
            elif row["decision"] == Decision.UNREACHABLE.value:
                if not bool(update.get("frontier_complete")):
                    raise ValueError(f"accepted unreachable oracle {label_id} requires a complete reviewed frontier")
            else:
                raise ValueError(f"accepted oracle {label_id} must be reachable or unreachable")
        row.update({
            "review_status": status, "review_reason": str(update.get("evidence") or ""),
            "reviewer": str(update.get("reviewer") or ""), "reviewed_at": str(update.get("reviewed_at") or ""),
        })
        accepted += status == "accepted"
        rejected += status == "rejected"
    write_parquet(dataset_dir / "oracle_labels.parquet", list(labels.values()))
    return {"updated": len(updates), "accepted": accepted, "rejected": rejected}
